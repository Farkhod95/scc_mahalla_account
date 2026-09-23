"""
Mahalla passporti (Excel) dan ma'lumotni bazaga yuklaydi.

    python manage.py import_mahalla_data files/data.xlsx --mahalla-code 1726277001
    python manage.py import_mahalla_data files/data.xlsx --mahalla-code 1726277001 --dry-run

Fayl kirill yozuvida bo'ladi — barcha matn lotinga o'giriladi
(`restapp.utils.translit`). Excel ichidagi rasmlar (PNG) drawing anchor'lari
bo'yicha qatorga bog'lanib, tegishli yozuvning `avatar` iga yoziladi.

Varaqlar va ular qayerga tushishi:

    умумий (faollar bloki)      -> Employee + Position
    умумий (sonlar bloki)       -> MahallaInformationCategory + MahallaInformation
    умумий (obyektlar bloki)    -> Object (ObjectCategory bo'yicha)
    рухий касаллар             -> MFYCitizen  category=ruhiy
    наркоманлар                -> MFYCitizen  category=narko
    жанжалкаш оилалар          -> MFYCitizen  category=janjal
    содир этилган жиноят        -> MahallaCrime + CrimeCategory
    ОЧ, АТИ                     -> o'tkazib yuboriladi (mos model yo'q)

Obyektlarning o'zi `умумий` dan olinadi (u yerda tashkilot nomi ham bor),
rasmlari esa alohida varaqlardan rahbar F.I.Sh bo'yicha topib qo'shiladi.
"""
import re
import zipfile
from datetime import datetime
from math import floor
from xml.etree import ElementTree as ET

from django.core.files.base import ContentFile
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from directory.models import Mahalla, Position
from monitoring.models import (CrimeCategory, Employee, MahallaCrime, MahallaInformation,
                               MahallaInformationCategory, MFYCitizen, Object, ObjectCategory)
from restapp.utils.translit import to_latin, to_latin_name

# ---------------------------------------------------------------- varaq nomlari
SHEET_UMUMIY = 'умумий'
# Bitta varaq turli fayllarda turlicha nomlanadi ('наркоманлар' / 'наркалогия'),
# shuning uchun har bir toifa uchun bir nechta nom qabul qilinadi.
SHEET_CITIZENS = {
    ('рухий касаллар', 'руҳий касаллар'): MFYCitizen.CATEGORY.RUHIY,
    ('наркоманлар', 'наркалогия', 'наркология'): MFYCitizen.CATEGORY.NARKO,
}
SHEET_JANJAL = 'жанжалкаш оилалар'
SHEET_CRIME = 'содир этилган жиноят'
SHEETS_SKIPPED = ('ОЧ', 'АТИ')

# `умумий` obyektlar blokidagi turkum nomi -> ObjectCategory.key.
# Kalitlar kichik harfda, bo'sh joylar siqilgan (norm_key).
OBJECT_CATEGORY_KEY = {
    'мактаб': 'maktab',
    'мактабгача таълим ташкилоти': 'bogcha',
    'оилавий поликлиника': 'poliklinika',
    'шифохона': 'shifoxona',
    'хусусий шифохона': 'xususiy_shifoxona',
    'олий таълим ташкилоти': 'oliy_talim',
    'коллеж': 'kollej_litsey',
    'колледж,литцей': 'kollej_litsey',
    'колледж,лицей': 'kollej_litsey',
    'ўқув маркази': 'oquv_markazi',
    'болалар майдончаси': 'bolalar_maydoni',
    'спорт майдончаси': 'sport_maydoni',
    'сартарошхона': 'sartaroshxona',
    'гўзаллик салони': 'gozallik_saloni',
    'дорихона': 'dorixona',
    'умумий овқатланиш': 'ovqatlanish_joyi',
    'тадбиркорлик субъектлари': 'tadbirkorlik',
    'савдо дўконлари': 'savdo_dokoni',
    'савдо мажмуаси': 'savdo_dokoni',
    'хостел': 'hostel',
    'футбол майдони': 'stadion',
    'талабалар турар жойи': 'talabalar_turar_joyi',
    'ҳаммом (сауна)': 'xamom(sauna)',
    'хамом (сауна)': 'xamom(sauna)',
    'банк': 'bank',
    'маданият ва истирохат боғлари': 'madiniyat_va_istirohat',
    'масжидлар': 'masjid',
    'масжид': 'masjid',
    'тўйхоналар': 'toyxona',
    'тўйхона': 'toyxona',
    'бозор': 'bozor',
    'музей': 'muzey',
    'мехмонхона': 'mexmonxona',
    # ko'plik/birlik va boshqa yozilish variantlari
    'шифохоналар': 'shifoxona',
    'мактабгача таълим': 'bogcha',
    'мактаб ташкилоти': 'maktab',
    'давлат шифохонаси': 'shifoxona',
    'лицей': 'kollej_litsey',
    'савдо дўкони': 'savdo_dokoni',
    'савдо дўконлар': 'savdo_dokoni',
    # lotincha yozilgan, transliteratsiya bilan mos tushmaydigan variantlar
    # ("субъектлари" -> "subektlari", faylda esa "subyektlari")
    'tadbirkorlik subyektlari': 'tadbirkorlik',

    # --- Shayxontohur (1-Gom) passportlaridan qo'shilgan variantlar ---
    # `ў`/`ғ` siz yozilgan imlo — fold_latin bilan ham mos tushmaydi
    # ("гузаллик" -> guzallik, "гўзаллик" -> gozallik).
    'гузаллик салони': 'gozallik_saloni',
    'туйхона': 'toyxona',
    'гўш дўкон': 'savdo_dokoni',
    'поликлиника': 'poliklinika',
    'жомеъ масжид': 'masjid',
    'мактаб ва хусусий мактаб': 'maktab',
    "o'rta umum talim maktabi": 'maktab',
    'хусусий бохча': 'bogcha',
    'otm': 'oliy_talim',
    'тадбиркорлик субъекти': 'tadbirkorlik',
    'тадбиркор': 'tadbirkorlik',
    'савдо дуконлари': 'savdo_dokoni',
    # Turkum o'rniga muassasaning O'Z NOMI yozilgan qatorlar
    'арзон хостел': 'hostel',
    'хотел': 'mexmonxona',
    'самарқанд дарвоза хотел': 'mexmonxona',
}

# `умумий` obyektlar blokida uchraydigan, lekin OBYEKT BO'LMAGAN qatorlar:
# fuqaro/jinoyat sonlari. Ular o'z varag'idan import qilinadi, shuning uchun
# bu yerda jimgina o'tkaziladi (ogohlantirish chiqarilmaydi).
NOT_AN_OBJECT_ROW = {
    'жанжалкаш оилалар',
    'содир этилган жиноят',
    'содир етилган жиноятлар',
    'наркологик диспансерда рўйчатда турадиган фуқоролар',
    'рухий касаллар диспансерида рўйхатда турадиган фуқоролар',
    'рухий касалликлар лиспансерида турувчи фуқаролар',
    'наркологик диспансерда рўйхатда турадиган фуқоролар',
    'спиртли ичимликка ружу қуйган шахлар',
    'рухий касаллар',
    'жанжалкаш оила',
    'гиёҳванд',
    'оч', 'ати',
}

# Alohida varaq nomi -> ObjectCategory.key (rasm va `умумий` da yo'q obyektlar uchun)
PHOTO_SHEET_CATEGORY = {
    'коллеж': 'kollej_litsey',
    'Колледж,лицей': 'kollej_litsey',
    'Мактаб': 'maktab',
    'Мактабгача таълим': 'bogcha',
    'Оилавий поликлиника': 'poliklinika',
    'Шифохона': 'shifoxona',
    'Хусусий шифохона': 'xususiy_shifoxona',
    'Олий таълим': 'oliy_talim',
    'Ўқув маркази': 'oquv_markazi',
    'Ўқув марказ': 'oquv_markazi',
    'Болалар майдончаси': 'bolalar_maydoni',
    'спорт майдончаси': 'sport_maydoni',
    'Сартарошхона': 'sartaroshxona',
    'Гўзаллик салони': 'gozallik_saloni',
    'Дорихона': 'dorixona',
    'Умумий овқатланиш': 'ovqatlanish_joyi',
    'тадбиркорлик субъекти': 'tadbirkorlik',
    'Савдо дуконлари': 'savdo_dokoni',
    'Савдо дўконлар': 'savdo_dokoni',
    'Савдо дўкони': 'savdo_dokoni',
    'Савдо дўкон': 'savdo_dokoni',
    'Савдо мажмуаси': 'savdo_dokoni',
    'хостел': 'hostel',
    'Футбол майдони': 'stadion',
    'талабалар турар жойи': 'talabalar_turar_joyi',
    'хамом (сауна)': 'xamom(sauna)',
    'банк': 'bank',
    'маданият ва истирохат боғлари,': 'madiniyat_va_istirohat',
    'масжид': 'masjid',
    'тўйхона': 'toyxona',
    'бозор': 'bozor',
    'музей': 'muzey',
    'мехмонхона': 'mexmonxona',
}

# Bazada mos ObjectCategory yo'q — bu varaqlar/turkumlar o'tkazib yuboriladi.
# Kerak bo'lsa avval ObjectCategory qo'shiladi, keyin shu ro'yxatdan olinadi.
CATEGORY_NOT_IN_DB = {
    'кутубхона', 'зиёратгоҳ', 'зиёратгох', 'қабристон',
    'маданият мерос объектлари',
}

# Fayldagi lavozim yozuvi -> bazadagi KANONIK nom.
# Kalit: lotinlashtirilgan nomning normalizatsiyasi (kichik harf, faqat harf/raqam).
# Manba fayllarda imlo turlicha ("М.Ф.Й риси", "МФЙ раиси", qisqartma nuqtali) —
# shuning uchun yangi Position yaratishdan oldin shu jadval tekshiriladi.
# "katta"/"kichik" — bu UNVON, alohida lavozim emas: hammasi
# 'Profilaktika inspektori' ga bog'lanadi.
POSITION_ALIASES = {
    "mfyrisi": "Mahalla raisi",
    "mfyraisi": "Mahalla raisi",
    "mahallaraisi": "Mahalla raisi",
    "mfyrayisi": "Mahalla raisi",
    "mahallafuqarolaryiginiraisi": "Mahalla raisi",

    "maxallaraisi": "Mahalla raisi",
    "mahalla": "Mahalla raisi",
    "maxalla": "Mahalla raisi",

    "rais": "Mahalla raisi",

    "profilaktikainspektori": "Profilaktika inspektori",
    "profilaktikakattainspektori": "Profilaktika inspektori",
    "profilaktikakichikinspektori": "Profilaktika inspektori",
    "kattaprofilaktikainspektori": "Profilaktika inspektori",
    # "инспектори" o'rniga "иснпектори" — Shayxontohur passportlarida
    # uchraydigan imlo xatosi (Kamolon Darvoza fayli).
    "profilaktikaisnpektori": "Profilaktika inspektori",
    "profilaktikakattaisnpektori": "Profilaktika inspektori",
}

# `умумий` da faollar bloki tugaganini bildiruvchi sarlavha kataklari (B ustuni).
# Faollar bloki bilan sonlar bloki orasida bo'sh qator bo'lmasa, sarlavhaning
# o'zi "xodim" bo'lib bazaga tushib qolardi: ismi "Obyekt nomi", lavozimi
# "Soni" bo'lgan yozuv va shu nomdagi keraksiz Position.
BLOCK_HEADER_LABELS = {
    "объект номи", "объекти номи", "обьект номи", "объектлар номи",
    "объект номлари", "ташкилот номи", "т/р",
}

# `умумий` sonlar blokidagi ko'rsatkich nomi -> bazadagi KANONIK
# MahallaInformationCategory nomi. Kalit: lotinlashtirilgan nomning
# normalizatsiyasi (`norm_name`).
# Manba fayllarda bir xil ko'rsatkich turlicha yoziladi ("Ахоли сони" /
# "Аҳоли сони", "Кўп қаватли уйлар" / "… уйлар сони"). Alias bo'lmasa
# bazadagi 'Aholi soni' turib 'Axoli soni' degan dublikat paydo bo'ladi
# va dashboardda takror qatorlar chiqadi.
INFO_CATEGORY_ALIASES = {
    "aholisoni": "Aholi soni",
    "axolisoni": "Aholi soni",
    "umumiyyermaydoni": "Umumiy yer maydoni (ga)",
    "umumiyyermaydoniga": "Umumiy yer maydoni (ga)",
    "umumiyermaydoni": "Umumiy yer maydoni (ga)",
    "kopqavatliuylar": "Ko‘p qavatli uylar soni",
    "kopqavatliuylarsoni": "Ko‘p qavatli uylar soni",
    "kopqavatliuyjoylar": "Ko‘p qavatli uylar soni",
    "xovliuyjoy": "Xovli uy-joy",
    "hovliuyjoy": "Xovli uy-joy",
    "xovliuyjoylar": "Xovli uy-joy",
}

# Jinoyat turkumining imlo variantlari -> bazadagi KANONIK nom.
# Kalit: lotinlashtirilgan nomning normalizatsiyasi (`norm_name`).
# Manba fayllarda bir xil turkum turlicha yoziladi va har variant uchun
# alohida turkum paydo bo'lsa, jinoyat filtrida takror qatorlar chiqadi.
CRIME_CATEGORY_ALIASES = {
    "giyooxvandlik": "Giyoxvandlik",
    "qushmachilik": "Qoʻshmachilik",
    "oylaviymayshiyzoravonlik": "Oilaviy maishiy zoʻravonlik",
    "tanjaroaxati": "Tan jaroxati",
}

# Jinoyat varag'ining "тасниф" ustunida turkum nomi o'rniga butun voqea bayoni
# yozilgan fayllar bor (Suzukota, Katta Oqtepa, Samarqand Darvoza). Bunday
# matndan turkum yasalsa u `CrimeCategory.name` (varchar 255) ga sig'maydi va
# PostgreSQL da butun fayl importi uzilib qoladi. Eng uzun haqiqiy turkum
# nomi ~40 belgi, shuning uchun chegara kengroq olindi.
CRIME_CATEGORY_MAX = 80

# Lavozim oldida mahalla nomi turishi mumkin ("Дўстлик МФЙ раиси") —
# har bir mahalla uchun alohida alias yozmaslik uchun OXIRI bo'yicha qoida.
POSITION_SUFFIX_RULES = (
    (re.compile(r"(mfy|mahalla|maxalla)ra[iy]?si$"), "Mahalla raisi"),
    (re.compile(r"profilaktika(katta|kichik)?inspektori$"),
     "Profilaktika inspektori"),
)

# Lavozim/ism/telefon BO'SH ekanini bildiruvchi yozuv — qator o'tkaziladi.
# "йуқ" (ў o'rniga у) ham uchraydi: lotinda "yuq" bo'ladi.
VACANT_NAMES = {"vokant", "vakant", "bosh", "yoq", "yuq", "vb"}

NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pkg": "http://schemas.openxmlformats.org/package/2006/relationships",
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
}


# ---------------------------------------------------------------- yordamchilar
def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def fold_latin(value):
    """Yozuvdan qat'i nazar taqqoslash uchun: lotinga o'girib, faqat harf/raqam.

    Blok kirillcha bo'lsa ham ayrim qatorlar lotinda yoziladi ("Bozor",
    "Sartaroshxona", "Tadbirkorlik subyektlari"). Faqat kirillcha kalit bilan
    solishtirilganda bunday qatorlar jimgina tushib qolardi va obyekt bazaga
    umuman kirmasdi.
    """
    return re.sub(r"[^0-9a-z]", "", to_latin(clean(value)).lower())


OBJECT_CATEGORY_FOLDED = {fold_latin(k): v for k, v in OBJECT_CATEGORY_KEY.items()}
CATEGORY_NOT_IN_DB_FOLDED = {fold_latin(k) for k in CATEGORY_NOT_IN_DB}
NOT_AN_OBJECT_ROW_FOLDED = {fold_latin(k) for k in NOT_AN_OBJECT_ROW}


def split_lines(value):
    """
    Katakni QATORLARGA bo'ladi. `clean()` yangi qatorni bo'sh joyga aylantiradi,
    shuning uchun ism/lavozim kabi "matn + sana" kataklari uchun shu kerak.
    """
    if value is None:
        return []
    return [clean(part) for part in re.split(r"[\r\n]+", str(value)) if clean(part)]


# Sana + ixtiyoriy "й"/"йил" ("11,01,1977 й") — "й" ham birga olib tashlanadi,
# aks holda ismda yolg'iz "Y" bo'lib qoladi.
DATE_RE = re.compile(r"\d{1,2}[.,/]\d{1,2}[.,/]\d{4}(?:\s*й(?:ил)?)?", re.IGNORECASE)
# "1983 й", "1985 йил" — faqat yil; sana sifatida ishlatilmaydi, lekin
# matndan olib tashlanadi (aks holda lavozim nomiga yopishib qoladi).
YEAR_RE = re.compile(r"\b(19|20)\d{2}\s*й(ил)?\b")

# Oy nomlari — lavozim katagida "2022 йил сентябрь" kabi qo'shimcha uchraydi
MONTH_RE = re.compile(
    r"\b(январ|феврал|март|апрел|май|июн|июл|август|сентябр|октябр|ноябр|декабр)\w*\b",
    re.IGNORECASE)

# Lavozim katagida tayinlanish vaqti turli ko'rinishda yoziladi va DATE_RE /
# YEAR_RE / MONTH_RE dan keyin ham qoldiq qolishi mumkin:
#   "МФЙ раиси 2025 йил май ойидан"        -> "май" olingach "ойидан" qoladi
#   "Профилактика катта инспектори 08.08."  -> chala sana (yili yozilmagan)
# Bunday qoldiq lavozim nomiga yopishib qolsa, POSITION_ALIASES bilan mos
# kelmaydi va bazada keraksiz yangi Position yaratiladi.
POSITION_NOISE_RE = re.compile(
    r"\b\d{1,2}\s*[.,/]\s*\d{1,2}\s*[.,/]?"   # chala sana: "08.08."
    r"|\bой(идан|дан|ида|и)?\b"                # oy nomi olingach qolgan "ойидан"
    r"|\bйил(идан|дан|и)?\b"                   # yolg'iz "йилдан"
    r"|\bдан\b"                                # "2020 йилдан" -> qolgan "дан"
    # Ba'zi fayllarda tayinlanish vaqti LOTINCHA yozilgan
    # ("Дўстлик МФЙ раиси 2025 yildan lavozimda") — kirillcha qoidalar
    # bunga tegmaydi va lavozim nomiga yopishib qoladi.
    r"|\byil(dan|da|gi|i)?\b"
    r"|\blavozim(da|ga|ida)?\b"
    r"|\b\d+\b",                               # boshqa yolg'iz raqamlar
    re.IGNORECASE)

# Ism katagida F.I.Sh dan keyin tug'ilgan sana/joy keladi:
#   "Шамсиева Норбуви Солиевна 31.03.1964 й.т. Жиззах шахар"
# Ism shu belgilardan OLDINGI qism deb olinadi.
NAME_CUT_RE = re.compile(
    r"\d{1,2}[.,/]\d{1,2}[.,/]\d{4}"      # 31.03.1964
    r"|\b(19|20)\d{2}\b"                   # 1971
    r"|\bй\.?\s*т\.?"                      # й.т.
    # Sana yozilmay, faqat tug'ilgan JOY qo'shilgan bo'lishi ham mumkin:
    # "Каримов Сардор Дамин ўғли Жиззах шаҳри" — yuqoridagi belgilar
    # bo'lmagani uchun joy nomi ismga yopishib qolardi.
    # Joy nomining O'ZI ham kesilishi kerak ("Жиззах шаҳри" — ikkalasi),
    # shuning uchun oldidagi so'z ham qamrab olinadi.
    r"|\s\S+\s+ша[ҳх]ри?\b"
    r"|\s\S+\s+ша[ҳх]ар\b"
    r"|\s\S+\s+туман(и|ida)?\b"
    r"|\s\S+\s+вилояти?\b",
    re.IGNORECASE
)


def take_date_out(parts):
    """
    Qatorlar ro'yxatidan sanani ajratib oladi -> (qolgan matn, sana, yil).

    Sana alohida qatorda ham ("Lavozim\n22.01.1992"), matn oxirida bo'shliqlar
    bilan ham ("Isayev Niyatilla    05.01.1965") uchraydi. Ba'zi fayllarda esa
    faqat yil yoziladi ("М.Ф.Й риси 1983 й") — u alohida qaytariladi.
    """
    rest, found, year = [], None, None
    for part in parts:
        if re.fullmatch(r"[\d.,/ й]+", part):          # butun qator — sana
            found = found or parse_date(part)
            continue
        m = DATE_RE.search(part)                        # matn ichidagi to'liq sana
        if m:
            found = found or parse_date(m.group(0))
            part = clean(part[:m.start()] + part[m.end():])
        m = YEAR_RE.search(part)                        # faqat yil
        if m:
            year = year or int(m.group(0)[:4])
            part = clean(part[:m.start()] + part[m.end():])
        part = clean(part.strip(" /-,"))
        if part:
            rest.append(part)
    return rest, found, year


def person_name(value):
    """
    Katakdan F.I.Sh ni ajratib, lotinga o'giradi.

    Ismdan keyin tug'ilgan sana va joy kelishi mumkin — bir qatorda ham
    ("Шамсиева Норбуви Солиевна 31.03.1964 й.т. Жиззах шахар"), yangi
    qatorda ham. Ism shu belgilardan oldingi qism deb olinadi, aks holda
    "Shamsiyeva … Jizzax shaxar" bo'lib qoladi va `умумий` bilan alohida
    varaqdagi yozuv mos kelmay, rasm bog'lanmaydi.
    """
    # Familiya va ism alohida qatorlarda bo'lishi mumkin
    # ("Хакимов\nАдхам Уринбосарович\n1958 йил") — avval birlashtiramiz,
    # keyin birinchi sana/yil belgisidan kesamiz.
    text = " ".join(split_lines(value))
    if not text:
        return ""
    m = NAME_CUT_RE.search(text)
    if m:
        text = text[:m.start()]
    return to_latin_name(clean(text).strip(" ,.-"))


def position_name(value):
    """
    Lavozim katagidan nomni ajratadi -> (nom, sana).

    Barcha qatorlar birlashtiriladi ("Маҳалла\\nраиси" -> "Mahalla raisi"),
    sana / yil / oy nomi va ulardan qolgan qo'shimchalar olib tashlanadi
    ("Профилактика катта инспектори 2022 йил сентябрь" ->
    "Profilaktika katta inspektori", "МФЙ раиси 2025 йил май ойидан" ->
    "MFY raisi").
    """
    text = " ".join(split_lines(value))
    if not text:
        return None, None

    appointed = None
    m = DATE_RE.search(text)
    if m:
        appointed = parse_date(m.group(0))
    text = DATE_RE.sub(" ", text)
    text = YEAR_RE.sub(" ", text)
    text = MONTH_RE.sub(" ", text)
    text = POSITION_NOISE_RE.sub(" ", text)
    text = clean(text).strip(" ,.-/")
    return (to_latin(text) or None), appointed


def parse_count(value):
    r"""
    Sonlar blokidagi katakdan sonni oladi: `'2540 та'` -> 2540, `'156 Г'` -> 156,
    `'50,2 (га)'` -> 50, `'111,5 г'` -> 112. Son topilmasa None.

    Ilgari barcha raqam bo'lmagan belgilar olib tashlanardi (`re.sub(r"\D", ...)`),
    shu sabab kasrli yer maydoni o'n barobar oshib ketardi: `'50,2 (га)'` -> 502.
    """
    m = re.search(r"\d+(?:[.,]\d+)?", clean(value))
    if not m:
        return None
    return floor(float(m.group(0).replace(",", ".")) + 0.5)


def norm_key(value):
    return clean(value).lower().rstrip(" .")


def resolve_sheet(wb, wanted):
    """
    Varaq nomini normalizatsiya bilan topadi -> haqiqiy nom yoki None.

    Manba fayllarda bir xil varaq turlicha yoziladi: katta/kichik harf
    ("Савдо дўконлар" / "савдо дўконлар"), oxirida ortiqcha probel
    ("маданият ва истирохат боғлари,  "). Aniq taqqoslashda bunday varaqlar
    jimgina tushib qolardi va obyektlari bazaga umuman kirmasdi.
    """
    if wanted in wb.sheetnames:
        return wanted
    target = norm_key(wanted)
    for name in wb.sheetnames:
        if norm_key(name) == target:
            return name
    return None


def norm_name(value):
    """Taqqoslash uchun: faqat harf va raqam, kichik harfda ('M.F.Y risi' -> 'mfyrisi')."""
    return re.sub(r"[^0-9a-z]", "", clean(value).lower())


# Sarlavha ustunlarini nomi bo'yicha topish uchun kalit so'zlar.
# Fayllarda ustun tartibi va sarlavha qatori har xil — shuning uchun
# qat'iy raqam emas, matn bo'yicha qidiriladi.
CITIZEN_HEADER = {
    "full_name": ("ф.и.ш", "фиш"),
    "surname": ("фамилия",),
    "given": ("исми",),
    "patronymic": ("шарифи",),
    "address": ("доимий яшаш", "яшаш манзил"),
    "account_type": ("исоб тоифаси",),
    "degree": ("даража",),
    "kind": ("тури",),
    "phone": ("телефон",),
    "photo": ("сурат",),
    "location": ("локац", "лакатц"),
}

# `умумий` faollar bloki. Odatdagi tartib: 2=Ф.И.Ш, 3=Лавозими, 5=манzil,
# 6=телефон — lekin ba'zi fayllarda ortiqcha "Tug'ilgan sanasi" ustuni bor va
# undan keyingi hamma narsa bir katak suriladi. Qat'iy raqamga ishonilsa
# manzil telefon maydoniga, tug'ilgan sana esa manzil maydoniga tushib qoladi.
#
# "туғилган" yolg'iz kalit sifatida YARAMAYDI: F.I.Sh sarlavhasining o'zi
# "Ф.И.Ш, туғилган йили ва жойи" deb yoziladi. Shuning uchun "сана" bilan.
ACTIVIST_HEADER = {
    "full_name": ("ф.и.ш", "фиш"),
    "position": ("лавозим",),
    "photo": ("расм",),
    "birthday": ("туғилган сана", "тугилган сана", "ilgan sanasi"),
    "address": ("яшаш манзил", "манзил"),
    "phone": ("телефон",),
}

OBJECT_HEADER = {
    "leader": ("раҳбар", "рахбар"),
    "photo": ("расм",),
    "phone": ("телефон",),
    "address": ("манзил",),
    "location": ("локац", "лакатц"),
}

# `умумий` obyektlar bloki. Odatdagi tartib: 2=Объект номи, 4=Ташкилот,
# 5=Раҳбар Ф.И.Ш, 6=Телефон, 7=манzil, 8=локация. Faollar blokidagi kabi,
# ba'zi fayllarda ortiqcha "Tug'ilgan sanasi" ustuni bor va undan keyingi
# ustunlar suriladi — o'shanda telefon maydoniga tug'ilgan sana tushardi.
UMUMIY_OBJECT_HEADER = {
    "label": ("объект номи", "обьект номи"),
    "org": ("ташкилот",),
    "leader": ("раҳбар", "рахбар"),
    "phone": ("телефон",),
    "address": ("манзил",),
    "location": ("локац", "лакатц"),
}


def find_header(ws, keywords, max_row=8):
    """
    Sarlavha qatorini va {kalit: ustun} moslikni topadi.

    Birinchi `max_row` qator ichidan eng ko'p kalit so'z mos kelganini oladi.
    Topilmasa (0, {}) qaytadi — chaqiruvchi qanday davom etishni o'zi hal qiladi.
    """
    # Ba'zi varaqlarda max_column 16384 bo'lib ketadi (bo'sh ustunlar) —
    # sarlavha har doim boshida, shuning uchun cheklaymiz.
    max_col = min(ws.max_column, 60)
    best_row, best = 0, {}
    for row in range(1, min(ws.max_row, max_row) + 1):
        found = {}
        for col in range(1, max_col + 1):
            label = norm_key(ws.cell(row, col).value)
            if not label:
                continue
            for key, words in keywords.items():
                if key not in found and any(w in label for w in words):
                    found[key] = col
        if len(found) > len(best):
            best_row, best = row, found
    return (best_row, best) if len(best) >= 2 else (0, {})


def parse_phone(value):
    """
    `91-198-58-24` -> `+998911985824`.

    Telefon katagida har doim ham telefon turmaydi: "йўқ"/"Yo'q" (= yo'q),
    bitta katakda ikkita raqam, ba'zan manzil yoki sana. Yo'qlik belgisi
    bo'lsa None, tanib bo'lmagan matn esa LOTINGA o'girib qaytariladi —
    aks holda bazada kirill matn qolib ketadi.
    """
    raw = clean(value)
    if not raw:
        return None
    if norm_name(to_latin(raw)) in VACANT_NAMES:
        return None

    digits = re.sub(r"\D", "", raw)
    if len(digits) == 9:
        return "+998" + digits
    if len(digits) == 12 and digits.startswith("998"):
        return "+" + digits
    # Bitta katakka ikkita raqam yozilgan ("95 001 00 44 93 517 13 13") —
    # birinchisi olinadi, modelda bitta maydon bor.
    if len(digits) == 18:
        return "+998" + digits[:9]
    if len(digits) == 24 and digits.startswith("998"):
        return "+" + digits[:12]
    return to_latin(raw) or None


def parse_date(value):
    """`17.02.1989`, `08,06,2026`, `2026 йил 20 март`, datetime -> date. Bo'lmasa None."""
    if isinstance(value, datetime):
        return value.date()
    raw = clean(value).replace(",", ".").replace("/", ".").replace(" й", "").strip(" .")
    m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", raw)
    if m:
        try:
            return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1))).date()
        except ValueError:
            return None
    return parse_worded_date(value)


# Oy nomi -> raqami. Jinoyat sanasi ba'zi fayllarda so'z bilan yoziladi:
# "2026 йил 20 март" (data-10), "20 март 2026 йил".
MONTH_NUM = {
    "январ": 1, "феврал": 2, "март": 3, "апрел": 4, "май": 5, "июн": 6,
    "июл": 7, "август": 8, "сентябр": 9, "октябр": 10, "ноябр": 11, "декабр": 12,
}


def parse_worded_date(value):
    """`2026 йил 20 март` / `20 март 2026 йил` -> date. Bo'lmasa None."""
    raw = clean(value).lower()
    if not raw:
        return None

    month = next((num for stem, num in MONTH_NUM.items() if stem in raw), None)
    if month is None:
        return None

    year = next((int(m.group(0)) for m in re.finditer(r"\b(19|20)\d{2}\b", raw)), None)
    if year is None:
        return None

    # Yildan boshqa 1-2 xonali son — kun. Bo'lmasa oyning 1-sanasi.
    day = next((int(d) for d in re.findall(r"\b\d{1,2}\b", raw)), 1)
    try:
        return datetime(year, month, day).date()
    except ValueError:
        return None


def parse_location(value):
    """
    `40,10454 67,86338` / `40.104433, 67.859765` -> (lat, lng) matn ko'rinishida.

    Nuqtasi tushib qolgan qiymat (`6785957`) tiklanadi, lekin faqat natija
    mahalla chegarasiga tushsa — aks holda None qaytadi.
    """
    raw = clean(value)
    if not raw:
        return None, None, None

    nums = re.findall(r"\d+(?:[.,]\d+)?", raw)
    if len(nums) < 2:
        return None, None, f"koordinata o'qilmadi: {raw!r}"

    def to_float(s):
        return float(s.replace(",", "."))

    lat, lng = to_float(nums[0]), to_float(nums[1])
    fixed = []

    # Nuqtasiz butun son: 6785957 -> 67.85957
    def restore(v):
        s = str(int(v))
        return float(s[:2] + "." + s[2:]) if len(s) > 3 else v

    if abs(lng) > 180:
        lng = restore(lng)
        fixed.append("longitude")
    if abs(lat) > 90:
        lat = restore(lat)
        fixed.append("latitude")

    note = None
    if fixed:
        note = (f"{' va '.join(fixed)} nuqtasi tushib qolgan deb tiklandi: "
                f"{raw!r} -> {lat}, {lng}")

    return lat, lng, note


def in_bbox(lat, lng, bbox, margin=0.05):
    """bbox = [min_lng, min_lat, max_lng, max_lat]."""
    if not bbox:
        return True
    min_lng, min_lat, max_lng, max_lat = bbox
    return (min_lng - margin <= lng <= max_lng + margin
            and min_lat - margin <= lat <= max_lat + margin)


# ------------------------------------------------------------------ rasmlar
def apply_src_rect(data, crop, label, warnings):
    """
    Excel'dagi kesishni (`a:srcRect`) rasmga qo'llaydi -> (baytlar, kesildimi).

    l/t/r/b — asl rasmning har tomonidan OLIB TASHLANADIGAN ulush, foizning
    mingdan bir ulushida (9369 = 9.369%). Excel faqat ko'rsatishda kesadi,
    fayl ichida asl rasm to'liq qoladi — shuning uchun qo'lda kesish kerak:
    bazaga Excel'da KO'RINIB turgan qism tushishi kerak.
    """
    if not crop:
        return data, False

    from io import BytesIO

    from PIL import Image

    def frac(key):
        return int(crop.get(key, 0) or 0) / 100000.0

    left, top, right, bottom = frac("l"), frac("t"), frac("r"), frac("b")
    # `<a:srcRect l="0" t="0" r="0" b="0"/>` — kesish yo'q. Qayta kodlash
    # rasmni yaxshilamaydi, faqat JPEG ni PNG ga aylantirib kattalashtiradi.
    if not any((left, top, right, bottom)):
        return data, False
    if left + right >= 1 or top + bottom >= 1:
        warnings.append(f"{label}: kesish chegarasi noto'g'ri {crop} — asl rasm olindi")
        return data, False

    try:
        with Image.open(BytesIO(data)) as im:
            w, h = im.size
            box = (round(w * left), round(h * top),
                   round(w * (1 - right)), round(h * (1 - bottom)))
            if box[2] <= box[0] or box[3] <= box[1]:
                warnings.append(f"{label}: kesishdan keyin bo'sh rasm — asl olindi")
                return data, False
            out = BytesIO()
            im.crop(box).save(out, format="PNG")
            return out.getvalue(), True
    except Exception as exc:
        warnings.append(f"{label}: rasm kesilmadi ({type(exc).__name__}: {exc})")
        return data, False


def load_photos(path):
    """{(varaq_nomi, qator): (fayl_nomi, baytlar)} — drawing anchor'lar bo'yicha."""
    photos, conflicts = {}, []

    with zipfile.ZipFile(path) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        wb_rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        rid_target = {r.get("Id"): r.get("Target")
                      for r in wb_rels.findall("pkg:Relationship", NS)}

        for sh in wb.find("main:sheets", NS):
            name = sh.get("name")
            target = rid_target[sh.get(f"{{{NS['r']}}}id")].lstrip("/")
            if not target.startswith("xl/"):
                target = "xl/" + target

            rels_path = target.replace("worksheets/", "worksheets/_rels/") + ".rels"
            if rels_path not in z.namelist():
                continue
            rels = ET.fromstring(z.read(rels_path))
            drawings = [r.get("Target") for r in rels.findall("pkg:Relationship", NS)
                        if "drawing" in (r.get("Type") or "")]
            if not drawings:
                continue

            dpath = "xl/" + drawings[0].replace("../", "")
            drels = ET.fromstring(
                z.read(dpath.replace("drawings/", "drawings/_rels/") + ".rels"))
            embed_media = {r.get("Id"): r.get("Target").replace("../", "xl/")
                           for r in drels.findall("pkg:Relationship", NS)}

            d = ET.fromstring(z.read(dpath))
            for tag in ("xdr:twoCellAnchor", "xdr:oneCellAnchor"):
                for a in d.findall(tag, NS):
                    frm = a.find("xdr:from", NS)
                    if frm is None:
                        continue
                    row = int(frm.find("xdr:row", NS).text) + 1
                    blip = a.find(".//a:blip", NS)
                    if blip is None:
                        continue
                    media = embed_media.get(blip.get(f"{{{NS['r']}}}embed"))
                    if not media or media not in z.namelist():
                        continue
                    key = (name, row)
                    if key in photos:
                        conflicts.append(f"{name} r{row}: bir katakda 2 rasm, "
                                         f"birinchisi olindi")
                        continue

                    src = a.find(".//a:srcRect", NS)
                    crop = dict(src.attrib) if src is not None and src.attrib else None
                    data = z.read(media)
                    filename = media.rsplit("/", 1)[-1]
                    if crop:
                        before = len(data)
                        data, cropped = apply_src_rect(
                            data, crop, f"{name} r{row}", conflicts)
                        if cropped:
                            # Kesilgan rasm PNG bo'lib qayta kodlanadi — kengaytma
                            # ham mos bo'lishi kerak, aks holda `image3.jpeg` nomli
                            # faylga nginx `image/jpeg` Content-Type beradi.
                            filename = filename.rsplit(".", 1)[0] + ".png"
                            conflicts.append(
                                f"{name} r{row}: Excel'dagi kesish qo'llandi "
                                f"({_crop_text(crop)}), {before // 1024} KB -> "
                                f"{len(data) // 1024} KB")

                    photos[key] = (filename, data)

    return photos, conflicts


def _crop_text(crop):
    return ", ".join(f"{k}={int(v) / 1000:.1f}%" for k, v in sorted(crop.items()) if v)


class Command(BaseCommand):
    """
    Passport faylini yuklaydigan asosiy buyruq.

    Meros oluvchi buyruq faqat `xlsx` bilan `mahalla_code` ni berib, bitta
    faylga bog'langan buyruq yaratishi mumkin — parser bitta joyda turadi.
    Ko'p fayl uchun `import_mahalla_folder` ishlatiladi.
    """

    help = "Mahalla passporti Excel faylini bazaga yuklaydi (kirill -> lotin)."

    # Meros oluvchi buyruqlar shu ikkitasini belgilaydi
    xlsx = None
    mahalla_code = None

    def add_arguments(self, parser):
        if self.xlsx is None:
            parser.add_argument("xlsx", help="Excel fayl yo'li")
        if self.mahalla_code is None:
            parser.add_argument("--mahalla-code", required=True,
                                help="Mahalla kodi, masalan 1726277001")
        parser.add_argument("--dry-run", action="store_true",
                            help="Bazaga yozmaydi, faqat nima bo'lishini ko'rsatadi")
        parser.add_argument("--no-photos", action="store_true",
                            help="Rasmlarni yuklamaydi")

    # ------------------------------------------------------------------ run
    def handle(self, *args, **options):
        path = self.xlsx or options["xlsx"]
        options["mahalla_code"] = self.mahalla_code or options.get("mahalla_code")
        self.dry = options["dry_run"]
        self.warnings = []
        self.stats = {}
        self.skipped_categories = set()

        try:
            self.mahalla = Mahalla.objects.select_related(
                "region", "district", "gom").get(code=options["mahalla_code"])
        except Mahalla.DoesNotExist:
            raise CommandError(f"Mahalla topilmadi: code={options['mahalla_code']}")

        # Bu loyihada hudud chegarasi (MahallaBoundary) modeli yo'q — bbox bo'sh
        # qoladi va koordinata chegara bo'yicha tekshirilmaydi. Chegara modeli
        # qo'shilsa, shu yerda o'zi ishlab ketadi.
        self.bbox = getattr(getattr(self.mahalla, "boundary", None), "bbox", None)

        self.stdout.write(self.style.MIGRATE_HEADING(
            f"Mahalla: {self.mahalla.name_uz} (id={self.mahalla.id}, "
            f"code={self.mahalla.code}) — {self.mahalla.district}"))
        if self.dry:
            self.stdout.write(self.style.WARNING("DRY-RUN: bazaga yozilmaydi\n"))

        try:
            wb = load_workbook(path, data_only=True)
        except Exception as exc:
            raise CommandError(f"Excel o'qilmadi: {exc}")

        self.photos = {} if options["no_photos"] else None
        if self.photos is None:
            self.photos, conflicts = load_photos(path)
            self.warnings.extend(conflicts)
            self.stdout.write(f"rasmlar topildi: {len(self.photos)}\n")

        try:
            with transaction.atomic():
                self.import_citizens(wb)
                self.import_janjal(wb)
                self.import_crimes(wb)
                self.import_umumiy(wb)
                if self.dry:
                    transaction.set_rollback(True)
        finally:
            wb.close()

        self.report()

    # -------------------------------------------------------------- saqlash
    def save_photo(self, obj, field, sheet, row):
        """
        Rasm bo'lsa `field` ga yozadi. dry-run da faqat sanaydi.

        Qayta import qilinganda eski fayl o'chiriladi — aks holda media/ da
        ishlatilmaydigan nusxalar to'planib qoladi.
        """
        found = self.photos.get((sheet, row))
        if not found:
            return False
        if not self.dry:
            filename, data = found
            existing = getattr(obj, field)
            if existing:
                existing.delete(save=False)
            existing.save(f"{obj.pk}_{filename}", ContentFile(data), save=True)
        self.stats["rasm"] = self.stats.get("rasm", 0) + 1
        return True

    def check_coords(self, lat, lng, where, note):
        if note:
            self.warnings.append(f"{where}: {note}")
        if lat is None:
            return None, None
        if not in_bbox(lat, lng, self.bbox):
            self.warnings.append(
                f"{where}: koordinata mahalla chegarasidan tashqarida "
                f"({lat}, {lng}) — bo'sh qoldirildi")
            return None, None
        return f"{lat}", f"{lng}"

    def bump(self, key, created):
        k = f"{key} (yangi)" if created else f"{key} (yangilandi)"
        self.stats[k] = self.stats.get(k, 0) + 1

    # ------------------------------------------------------------ fuqarolar
    def import_citizens(self, wb):
        """
        рухий касаллар / наркоманлар.

        Ikki xil ko'rinish uchraydi:
          A) sarlavhali jadval — Фамилия / Исми / Шарифи alohida ustunlarda
             (+ ЖШШИР, ҳужжат, ҳисоб тоифаси, ba'zan Лакатция)
          B) sarlavhasiz — A=№, B=to'liq F.I.Sh, C=tug'ilgan sana, D=manzil,
             E=tashxis kodi, F=guruh, G=vasiylik, H=mahalla, I=hisob sanasi,
             K=lokatsiya

        Qaysi ko'rinish ekani sarlavhadan aniqlanadi, taxmin qilinmaydi.
        """
        for names, category in SHEET_CITIZENS.items():
            sheet = next((s for s in map(lambda n: resolve_sheet(wb, n), names) if s), None)
            if sheet is None:
                self.warnings.append(f"varaq yo'q: {' / '.join(names)}")
                continue
            ws = wb[sheet]
            header_row, cols = find_header(ws, CITIZEN_HEADER)

            if cols.get("surname") or cols.get("full_name"):
                self._citizens_with_header(ws, sheet, category, header_row, cols)
            else:
                self.warnings.append(
                    f"{sheet}: sarlavha topilmadi — ustunsiz ko'rinish deb o'qildi "
                    f"(B=F.I.Sh, D=manzil, E/F=tashxis, K=lokatsiya)")
                self._citizens_plain(ws, sheet, category)

    def _citizen_save(self, sheet, row, category, full_name, address,
                      degree=None, ctype=None, location=None, phone=None):
        if not full_name:
            return
        lat, lng, note = parse_location(location)
        cx, cy = self.check_coords(lat, lng, f"{sheet} r{row}", note)

        obj, created = self.upsert_by_name(
            MFYCitizen, full_name, {"category": category},
            defaults={
                "address": address or None,
                "phone": phone or None,
                "degree": degree or None,
                "type": ctype or None,
                "coordinate_x": cx, "coordinate_y": cy,
                "region": self.mahalla.region,
                "district": self.mahalla.district,
                "gom": self.mahalla.gom,
                "is_active": True,
            })
        if obj is None:
            self.warnings.append(
                f"{sheet} r{row}: F.I.Sh da harf/raqam yo'q ({full_name!r}) — "
                f"fuqaro qo'shilmadi")
            return
        self.bump(f"MFYCitizen {category}", created)
        self.save_photo(obj, "avatar", sheet, row)

    def _citizens_with_header(self, ws, sheet, category, header_row, cols):
        """
        Sarlavhali jadval. Ism bitta ustunda (`Ф.И.Ш`) ham,
        uchta ustunda (`Фамилия`/`Исми`/`Шарифи`) ham bo'lishi mumkin.
        """
        def cell(row, key):
            col = cols.get(key)
            return clean(ws.cell(row, col).value) if col else ""

        for row in range(header_row + 1, ws.max_row + 1):
            if cols.get("full_name"):
                full_name = person_name(ws.cell(row, cols["full_name"]).value)
            else:
                parts = [cell(row, k) for k in ("surname", "given", "patronymic")]
                full_name = to_latin_name(" ".join(p for p in parts if p))
            if not full_name:
                continue

            # `Ҳисоб тоифаси` (ruhiy -> daraja, narko -> tur) yoki
            # alohida `даражаси` / `тури` ustunlari
            account = to_latin(cell(row, "account_type")) or None
            degree = to_latin(cell(row, "degree")) or None
            kind = to_latin(cell(row, "kind")) or None
            if category == MFYCitizen.CATEGORY.RUHIY:
                degree, kind = degree or account, kind
            else:
                degree, kind = degree, kind or account

            self._citizen_save(
                sheet, row, category, full_name,
                address=to_latin(cell(row, "address")),
                degree=degree, ctype=kind,
                phone=parse_phone(cell(row, "phone")),
                location=cell(row, "location"))

    def _citizens_plain(self, ws, sheet, category):
        """Sarlavhasiz ko'rinish: B=F.I.Sh, D=manzil, E=tashxis, F=guruh, K=lokatsiya."""
        for row in range(1, ws.max_row + 1):
            name = clean(ws.cell(row, 2).value)
            number = clean(ws.cell(row, 1).value)
            # ma'lumot qatori: A ustunida tartib raqami, B da kamida ikki so'zli ism
            if not name or not re.fullmatch(r"\d+", number) or len(name.split()) < 2:
                continue
            diagnosis = clean(ws.cell(row, 5).value)
            group = clean(ws.cell(row, 6).value)
            degree = ", ".join(to_latin(x) for x in (diagnosis, group) if x) or None
            self._citizen_save(
                sheet, row, category, to_latin_name(name),
                address=to_latin(clean(ws.cell(row, 4).value)),
                degree=degree if category == MFYCitizen.CATEGORY.RUHIY else None,
                ctype=None if category == MFYCitizen.CATEGORY.RUHIY else degree,
                location=clean(ws.cell(row, 11).value))

    def import_janjal(self, wb):
        """жанжалкаш оилалар — ustunlar sarlavha bo'yicha topiladi."""
        sheet = resolve_sheet(wb, SHEET_JANJAL)
        if sheet is None:
            self.warnings.append(f"varaq yo'q: {SHEET_JANJAL}")
            return
        ws = wb[sheet]
        header_row, cols = find_header(ws, CITIZEN_HEADER)
        if not (cols.get("full_name") or cols.get("surname")):
            self.warnings.append(f"{SHEET_JANJAL}: sarlavha topilmadi, o'tkazib yuborildi")
            return
        self._citizens_with_header(ws, SHEET_JANJAL, MFYCitizen.CATEGORY.JANJAL,
                                   header_row, cols)

    # -------------------------------------------------------------- jinoyat
    def import_crimes(self, wb):
        """содир этилган жиноят — sarlavha r3, ma'lumot r4 dan."""
        sheet = resolve_sheet(wb, SHEET_CRIME)
        if sheet is None:
            self.warnings.append(f"varaq yo'q: {SHEET_CRIME}")
            return
        ws = wb[sheet]
        for row in range(4, ws.max_row + 1):
            article = clean(ws.cell(row, 5).value)
            classification = clean(ws.cell(row, 4).value)
            if not article and not classification:
                continue

            category = None
            if classification:
                category = self.resolve_crime_category(to_latin(classification), row)

            lat, lng, note = parse_location(ws.cell(row, 6).value)
            cx, cy = self.check_coords(lat, lng, f"{SHEET_CRIME} r{row}", note)
            date = parse_date(ws.cell(row, 3).value)
            if date is None:
                self.warnings.append(
                    f"{SHEET_CRIME} r{row}: sana o'qilmadi "
                    f"({clean(ws.cell(row, 3).value)!r})")

            obj, created = MahallaCrime.objects.update_or_create(
                mahalla=self.mahalla, article=to_latin(article) or "-", date=date,
                defaults={
                    "category": category,
                    "description": to_latin(clean(ws.cell(row, 2).value)) or None,
                    "coordinate_x": cx, "coordinate_y": cy,
                    "region": self.mahalla.region, "district": self.mahalla.district,
                    "gom": self.mahalla.gom,
                })
            self.bump("MahallaCrime", created)

    def resolve_crime_category(self, name, row):
        """
        Jinoyat turkumini bazadagi MAVJUD yozuvga bog'laydi; faqat chinakam
        yo'q bo'lsa yangisini yaratadi.

        Manba fayllarda imlo bir xil emas ("Oʻgʻirlik" / "oʻgʻirlik"), aniq
        taqqoslashda esa har bir variant uchun alohida turkum paydo bo'lardi
        va jinoyat filtrida bitta turkum bir necha marta ko'rinardi.

        Tasnif o'rniga voqea bayoni yozilgan bo'lsa turkum bog'lanmaydi —
        jinoyatning o'zi baribir saqlanadi, turkumi admin paneldan qo'yiladi.
        """
        if len(name) > CRIME_CATEGORY_MAX:
            self.warnings.append(
                f"{SHEET_CRIME} r{row}: tasnif ustunida turkum emas, voqea "
                f"bayoni yozilgan ({len(name)} belgi) — turkum bog'lanmadi: "
                f"{name[:60]!r}…")
            return None

        # setdefault — bir xil normalizatsiyali turkumlardan ENG ESKISI
        # (eng kichik id) tanlanadi, chunki jinoyatlar odatda o'shanga bog'langan.
        existing = {}
        for category in CrimeCategory.objects.all():
            existing.setdefault(norm_name(category.name_uz or category.name), category)

        canonical = CRIME_CATEGORY_ALIASES.get(norm_name(name))
        target = canonical or name
        found = existing.get(norm_name(target))
        if found:
            if (found.name_uz or found.name) != name:
                self.warnings.append(
                    f"jinoyat turkumi '{name}' -> mavjud "
                    f"'{found.name_uz or found.name}' (id={found.pk}) ga bog'landi")
            return found

        # Alias bo'lsa KANONIK nom bilan yaratiladi: aks holda keyingi fayl
        # yana kanonik nom bo'yicha qidirib topa olmaydi va har bir faylda
        # bitta turkumdan yangisi paydo bo'lardi.
        category = CrimeCategory.objects.create(
            name=target, name_uz=target, name_ru=target, name_en=target)
        self.bump("CrimeCategory", True)
        return category

    # --------------------------------------------------------------- умумий
    def check_target_mahalla(self, ws):
        """
        Varaq sarlavhasidagi mahalla nomi tanlangan mahallaga mos kelishini
        tekshiradi.

        Bir xil nomli mahallalar bir nechta tumanda uchraydi (masalan "Yoshlik
        MFY" — 4 ta), shuning uchun noto'g'ri `--mahalla-code` bilan boshqa
        mahallaga yozib yuborish xavfi bor.
        """
        title = ""
        for row in range(1, 4):
            title = clean(ws.cell(row, 1).value)
            if title:
                break
        if not title:
            return

        latin = norm_name(to_latin(title))
        expected = norm_name((self.mahalla.name_uz or "").replace("MFY", ""))
        self.stdout.write(f"varaq sarlavhasi: {to_latin(title)[:90]}")
        if expected and expected not in latin:
            self.warnings.append(
                f"DIQQAT: varaq sarlavhasida '{to_latin(title)[:60]}' deyilgan, "
                f"tanlangan mahalla esa '{self.mahalla.name_uz}' "
                f"(code={self.mahalla.code}, {self.mahalla.district}) — "
                f"mahalla kodini tekshiring")

    def import_umumiy(self, wb):
        sheet = resolve_sheet(wb, SHEET_UMUMIY)
        if sheet is None:
            self.warnings.append(f"varaq yo'q: {SHEET_UMUMIY}")
            return
        ws = wb[sheet]
        self.check_target_mahalla(ws)
        self._umumiy_activists(ws)
        self._umumiy_counts(ws)
        self._umumiy_objects(ws)

    def _umumiy_activists(self, ws):
        """Faollar: sarlavha 'т/р | Ф.И.Ш' qatoridan keyin, bo'sh qatorgacha."""
        start = self._find_row(ws, lambda r: norm_key(ws.cell(r, 3).value).startswith(
            "лавозими"))
        if start is None:
            self.warnings.append("умумий: faollar sarlavhasi topilmadi")
            return

        # Ustunlar sarlavha matni bo'yicha aniqlanadi; topilmagani odatdagi
        # o'rnida deb olinadi (`ACTIVIST_HEADER` izohiga qarang).
        cols = {}
        for col in range(1, min(ws.max_column, 30) + 1):
            label = norm_key(ws.cell(start, col).value)
            if not label:
                continue
            for key, words in ACTIVIST_HEADER.items():
                if key not in cols and any(w in label for w in words):
                    cols[key] = col
        name_col = cols.get("full_name", 2)
        role_col = cols.get("position", 3)
        addr_col = cols.get("address", 5)
        phone_col = cols.get("phone", 6)

        empty_streak = 0
        for row in range(start + 1, ws.max_row + 1):
            # Keyingi blok sarlavhasi — faollar tugadi
            if norm_key(ws.cell(row, name_col).value) in BLOCK_HEADER_LABELS:
                break

            full_name = person_name(ws.cell(row, name_col).value)
            role, appointed = position_name(ws.cell(row, role_col).value)

            # "Вокант"/"ВБ" — lavozim bo'sh degani. Bunday qatordan xodim
            # yaratilsa bazada "Vokant" ismli soxta rais paydo bo'ladi.
            if full_name and norm_name(full_name) in VACANT_NAMES:
                self.warnings.append(
                    f"умумий r{row}: lavozim bo'sh ({full_name!r}) — "
                    f"xodim yaratilmadi")
                continue

            if not full_name:
                # Ism yo'q, lekin lavozim bor — sarlavha/bo'sh qator bo'lishi mumkin
                if role:
                    self.warnings.append(
                        f"умумий r{row}: lavozim '{role}' bor, lekin F.I.Sh bo'sh — "
                        f"qator o'tkazib yuborildi")
                empty_streak += 1
                if empty_streak >= 2:
                    break
                continue
            empty_streak = 0

            _, birthday, _ = take_date_out(split_lines(ws.cell(row, name_col).value))
            # Alohida "Tug'ilgan sanasi" ustuni bo'lsa — F.I.Sh katagida sana
            # yozilmagan bo'lishi mumkin.
            if birthday is None and cols.get("birthday"):
                birthday = parse_date(ws.cell(row, cols["birthday"]).value)
            position = self.resolve_position(role, row)

            phone = parse_phone(ws.cell(row, phone_col).value) or "-"
            obj, created = self.upsert_employee(
                full_name, phone, row,
                defaults={
                    "position": position,
                    "date_of_birthday": birthday,
                    "date_of_appointment": appointed,
                    "phone_number": phone,
                    "address": to_latin(clean(ws.cell(row, addr_col).value)) or None,
                    "type": "mahalla",
                    "region": self.mahalla.region, "district": self.mahalla.district,
                    "gom": self.mahalla.gom,
                })
            self.bump("Employee", created)
            # rasmlar haqiqiy varaq nomi bo'yicha indekslangan (' умумий' ham uchraydi)
            self.save_photo(obj, "avatar", ws.title, row)

    def upsert_by_name(self, model, full_name, extra, defaults, match_phone=None):
        """
        Yozuvni mahalla ichida NORMALIZATSIYALANGAN ism bo'yicha topib
        yangilaydi; topilmasa yangisini yaratadi -> (obyekt, yaratildimi).

        `update_or_create(full_name=...)` ism aynan mos kelishini talab qiladi.
        Bazadagi qo'lda kiritilgan yozuvlarda transliteratsiya boshqacha
        ("Xolboʻta Qizi" / "Xolbo'ta qizi") va qayta importda bir odam ikki
        marta paydo bo'lardi.

        Rahbari ko'rsatilmagan obyektlarda ism bo'sh bo'ladi — ular bir-biriga
        qo'shilib ketmasligi uchun avval `match_phone`, u ham bo'lmasa
        tashkilot nomi bo'yicha solishtiriladi. Uchalasi ham bo'sh bo'lsa
        yozuvni umuman aniqlab bo'lmaydi (faylda tasodifiy "\" kabi katak
        uchraydi) va u har importda yangi dublikat yaratardi — shuning uchun
        (None, False) qaytariladi va chaqiruvchi qatorni o'tkazib yuboradi.
        """
        qs = model.objects.filter(mahalla=self.mahalla, **extra)
        target = norm_name(full_name)
        org = norm_name(defaults.get("name") or "")
        if target:
            found = next((o for o in qs if norm_name(o.full_name) == target), None)
        elif match_phone:
            found = next((o for o in qs if o.phone_number == match_phone), None)
        elif org:
            found = next((o for o in qs
                          if norm_name(getattr(o, "name", "")) == org), None)
        else:
            return None, False
        if found is None:
            return model.objects.create(
                mahalla=self.mahalla, full_name=full_name, **extra, **defaults), True

        found.full_name = full_name
        self.apply_defaults(found, defaults)
        found.save()
        return found, False

    def upsert_employee(self, full_name, phone, row, defaults):
        """
        Xodimni mahalla ichida TELEFON yoki normalizatsiyalangan ism bo'yicha
        topib yangilaydi; topilmasa yangisini yaratadi -> (obyekt, yaratildimi).

        Ilgari `update_or_create(mahalla, full_name)` ishlatilardi — ism aynan
        mos kelishi kerak edi. Bazadagi yozuvlar qo'lda kiritilgani uchun
        transliteratsiya boshqacha ("Abdullaevna" / "Abdullayevna",
        "o'g'li" / "oʻgʻli") va qayta importda BIR ODAM IKKI MARTA paydo
        bo'lardi. Telefon raqami esa ikkalasida ham bir xil bo'ladi.
        """
        found = None
        if phone and phone != "-":
            matches = list(Employee.objects.filter(
                mahalla=self.mahalla, phone_number=phone).order_by("id"))
            if matches:
                found = matches[0]
                if len(matches) > 1:
                    self.warnings.append(
                        f"умумий r{row}: {phone} raqamli {len(matches)} ta xodim bor "
                        f"(id={[m.pk for m in matches]}) — birinchisi yangilandi")

        if found is None:
            target = norm_name(full_name)
            found = next((e for e in Employee.objects.filter(mahalla=self.mahalla)
                          if norm_name(e.full_name) == target), None)

        if found is None:
            return Employee.objects.create(
                mahalla=self.mahalla, full_name=full_name, **defaults), True

        if found.full_name != full_name:
            self.warnings.append(
                f"умумий r{row}: mavjud xodim '{found.full_name}' (id={found.pk}) "
                f"'{full_name}' ga yangilandi")
        found.full_name = full_name
        self.apply_defaults(found, defaults)
        found.save()
        return found, False

    @staticmethod
    def apply_defaults(obj, defaults):
        """
        Fayldagi qiymatlarni yozadi, lekin BO'SH qiymat bilan to'ldirilgan
        maydonni o'chirmaydi.

        Manba passportlar to'liq emas — masalan Shodlik faylida inspektorning
        lavozim katagi umuman bo'sh. Shunday faylni qayta import qilganda
        avval qo'lda to'g'rilangan lavozim `None` bilan yuvilib ketardi.
        """
        for field, value in defaults.items():
            if value in (None, "") and getattr(obj, field, None) not in (None, ""):
                continue
            setattr(obj, field, value)

    def resolve_position(self, name, row):
        """
        Lavozimni bazadagi MAVJUD yozuvga bog'laydi; faqat chinakam yo'q bo'lsa
        yangisini yaratadi.

        Manba fayllarda imlo turlicha ("М.Ф.Й риси" — "раиси" xato yozilgan),
        shuning uchun oldin POSITION_ALIASES, keyin mavjud nomlar bilan
        normalizatsiya qilib taqqoslanadi. Aks holda "Mahalla raisi" bor
        turib "M.F.Y risi" degan dublikat paydo bo'ladi va rais filtrga
        tushmay qoladi.
        """
        if not name:
            return None

        canonical = POSITION_ALIASES.get(norm_name(name))
        if canonical is None:
            for pattern, value in POSITION_SUFFIX_RULES:
                if pattern.search(norm_name(name)):
                    canonical = value
                    break
        target = canonical or name

        existing = {norm_name(p.name_uz or p.name): p for p in Position.objects.all()}
        found = existing.get(norm_name(target))
        if found:
            if canonical and norm_name(name) != norm_name(found.name_uz or found.name):
                self.warnings.append(
                    f"умумий r{row}: lavozim '{name}' -> mavjud "
                    f"'{found.name_uz or found.name}' (id={found.pk}) ga bog'landi")
            return found

        position = Position.objects.create(
            name=target, name_uz=target, name_ru=target, name_en=target)
        self.bump("Position", True)
        self.warnings.append(
            f"умумий r{row}: YANGI lavozim yaratildi '{target}' (id={position.pk}) — "
            f"bazada mos yozuv topilmadi")
        return position

    def resolve_info_category(self, name, row):
        """
        Ko'rsatkich turkumini bazadagi MAVJUD yozuvga bog'laydi; faqat chinakam
        yo'q bo'lsa yangisini yaratadi -> (turkum, yaratildimi).

        `Position` dagi kabi: manba fayllarda bir xil ko'rsatkich turlicha
        yoziladi ("Ахоли сони" / "Аҳоли сони"), aks holda bazada
        'Aholi soni' turib 'Axoli soni' degan dublikat paydo bo'ladi va
        dashboardda ikkita qator chiqadi.
        """
        canonical = INFO_CATEGORY_ALIASES.get(norm_name(name))
        target = canonical or name

        existing = {norm_name(c.name): c
                    for c in MahallaInformationCategory.objects.all()}
        found = existing.get(norm_name(target))
        if found:
            if canonical and norm_name(name) != norm_name(found.name):
                self.warnings.append(
                    f"умумий r{row}: ko'rsatkich '{name}' -> mavjud "
                    f"'{found.name}' (id={found.pk}) ga bog'landi")
            return found, False

        category = MahallaInformationCategory.objects.create(name=target)
        self.warnings.append(
            f"умумий r{row}: YANGI ko'rsatkich turkumi yaratildi '{target}' "
            f"(id={category.pk}) — bazada mos yozuv topilmadi")
        return category, True

    def _umumiy_counts(self, ws):
        """Sonlar bloki: sarlavha 'Объект номи | Сони | хонадонлар сони'."""
        start = self._find_row(ws, lambda r: norm_key(ws.cell(r, 4).value).startswith(
            "хонадонлар"))
        if start is None:
            # Ba'zi fayllarda 4-ustun bo'sh va son "Ер майдони гектар" (5-ustun)
            # da turadi — sarlavhada "хонадонлар" umuman yo'q. Bunday sarlavhani
            # obyektlar bloki sarlavhasidan ("Объект номи | Сони | Ташкилот …")
            # ajratuvchi yagona belgi — 4-ustunda "Ташкилот" yo'qligi.
            start = self._find_row(ws, lambda r: (
                norm_key(ws.cell(r, 2).value).startswith(("объект", "обьект"))
                and norm_key(ws.cell(r, 3).value).startswith("сони")
                and "ташкилот" not in norm_key(ws.cell(r, 4).value)))
        if start is None:
            self.warnings.append("умумий: sonlar bloki topilmadi")
            return

        seen = set()
        for row in range(start + 1, ws.max_row + 1):
            label = clean(ws.cell(row, 2).value)
            if not label:
                break

            # Odatda qiymat 3-ustunda ("Сони"). Ba'zi fayllarda esa u
            # "Ер майдони гектар" (5) yoki "хонадонлар сони" (4) ustuniga
            # yozilgan — "Умумий ер майдони" da deyarli har doim shunday.
            # Faqat 3-ustunga qarasak bunday qatorlar butunlay tushib qolardi.
            raw_count, source_col = clean(ws.cell(row, 3).value), 3
            if not raw_count:
                for col in (5, 4):
                    alt = clean(ws.cell(row, col).value)
                    if alt:
                        raw_count, source_col = alt, col
                        break
            if not raw_count:
                continue

            count = parse_count(raw_count)
            if count is None:
                self.warnings.append(f"умумий r{row}: son o'qilmadi ({raw_count!r})")
                continue
            if source_col != 3:
                self.warnings.append(
                    f"умумий r{row} '{clean(label)}': qiymat {source_col}-ustundan "
                    f"olindi ({raw_count!r} -> {count}), 3-ustun bo'sh edi")

            name = to_latin(label)
            category, created = self.resolve_info_category(name, row)
            if created:
                self.bump("MahallaInformationCategory", True)

            # Takror IMLO bo'yicha emas, TURKUM bo'yicha aniqlanadi: bitta
            # faylda "Ахоли сони" ikki marta ("Axoli soni" / "Aholi soni")
            # yozilgan bo'lsa ham ikkalasi bitta turkumga tushadi.
            if category.pk in seen:
                self.warnings.append(
                    f"умумий r{row}: '{name}' takrorlandi, birinchisi qoldirildi")
                continue
            seen.add(category.pk)

            extra = ""
            if source_col == 3:
                extra = clean(ws.cell(row, 4).value) or clean(ws.cell(row, 5).value)
            if extra:
                self.warnings.append(
                    f"умумий r{row} '{name}': qo'shimcha ustun qiymati ({extra!r}) "
                    f"modelda saqlanmaydi")

            _, created = MahallaInformation.objects.update_or_create(
                mahalla=self.mahalla, category=category,
                defaults={"count": count,
                          "region": self.mahalla.region,
                          "district": self.mahalla.district,
                          "gom": self.mahalla.gom})
            self.bump("MahallaInformation", created)

    def _umumiy_objects(self, ws):
        """Obyektlar bloki: sarlavha 'Ташкилот рақами, ёки номи' ustuni bo'yicha."""
        start = self._find_row(ws, lambda r: "ташкилот" in norm_key(ws.cell(r, 4).value))
        if start is None:
            self.warnings.append("умумий: obyektlar bloki topilmadi")
            return

        # Ustunlar sarlavha matni bo'yicha (`UMUMIY_OBJECT_HEADER` izohiga qarang)
        cols = {}
        for col in range(1, min(ws.max_column, 30) + 1):
            head = norm_key(ws.cell(start, col).value)
            if not head:
                continue
            for key, words in UMUMIY_OBJECT_HEADER.items():
                if key not in cols and any(w in head for w in words):
                    cols[key] = col
        label_col = cols.get("label", 2)
        org_col = cols.get("org", 4)
        leader_col = cols.get("leader", 5)
        phone_col = cols.get("phone", 6)
        addr_col = cols.get("address", 7)
        loc_col = cols.get("location", 8)

        photo_index = self._object_photo_index(ws)
        categories = {c.key: c for c in ObjectCategory.objects.all()}
        current_key = None
        seen_leaders = set()
        seen_phones = set()
        seen_rows = set()

        for row in range(start + 1, ws.max_row + 1):
            label = clean(ws.cell(row, label_col).value)
            unknown_label = None
            if label:
                normalized = norm_key(label)
                folded = fold_latin(label)
                key = OBJECT_CATEGORY_KEY.get(normalized) or OBJECT_CATEGORY_FOLDED.get(folded)
                if key:
                    current_key = key
                else:
                    # Bazada turkumi yo'q (kutubxona, qabriston…) yoki umuman
                    # obyekt emas (jinoyat/fuqaro sonlari) — ikkalasi ham o'tkaziladi
                    if normalized in CATEGORY_NOT_IN_DB or folded in CATEGORY_NOT_IN_DB_FOLDED:
                        self.skipped_categories.add(label)
                    elif folded not in NOT_AN_OBJECT_ROW_FOLDED:
                        unknown_label = label
                    current_key = None

            leader = clean(ws.cell(row, leader_col).value)
            if not leader:
                continue
            if current_key is None:
                # Rahbari bor, turkumi tanilmadi — jimgina yo'qotmaymiz
                if unknown_label:
                    self.warnings.append(
                        f"умумий r{row}: turkum tanilmadi ({unknown_label!r}) — "
                        f"obyekt qo'shilmadi")
                continue

            category = categories.get(current_key)
            if category is None:
                self.warnings.append(
                    f"умумий r{row}: ObjectCategory topilmadi (key={current_key})")
                continue

            full_name = person_name(ws.cell(row, leader_col).value)

            # Ba'zi passportlarda obyektlar jadvali ikki-uch marta ko'chirilgan
            # va NUSXADA turkum ustuni bo'sh qolgan. Bunday qator o'zidan
            # oldingi turkumni meros oladi va bitta obyekt xaritada ikki xil
            # turkumda, bir xil koordinatada ikki marta chiqadi (Zangiota
            # passportida 59 ta shunday qator bor edi).
            #
            # Turkumi YOZILGAN takror esa ataylab: bitta obyekt rostdan ikki
            # turkumga tegishli bo'lishi mumkin ("Sartaroshxona" + "Goʻzallik
            # saloni", "Maktab" + "Sport maydonchasi") — unga tegilmaydi.
            source_key = (norm_name(clean(ws.cell(row, org_col).value)),
                          norm_name(full_name))
            if not label and source_key in seen_rows:
                self.warnings.append(
                    f"умумий r{row}: jadvalning takroriy nusxasi (turkum ustuni "
                    f"bo'sh, '{current_key}' meros olingan) — obyekt qo'shilmadi")
                continue
            seen_rows.add(source_key)

            lat, lng, note = parse_location(ws.cell(row, loc_col).value)
            cx, cy = self.check_coords(lat, lng, f"умумий r{row}", note)

            # "Yo'q"/"Vokant" — rahbar yo'q degani, ism emas. Aks holda
            # obyekt rahbari sifatida bazaga "Yo'q" yozilib qoladi.
            if norm_name(full_name) in VACANT_NAMES:
                full_name = ""

            phone = parse_phone(ws.cell(row, phone_col).value) or "-"
            obj, created = self.upsert_by_name(
                Object, full_name, {"category": category},
                match_phone=phone if phone != "-" else None,
                defaults={
                    "name": to_latin(clean(ws.cell(row, org_col).value))[:100] or None,
                    "phone_number": phone,
                    "address": to_latin(clean(ws.cell(row, addr_col).value)) or None,
                    "coordinate_x": cx, "coordinate_y": cy,
                    "region": self.mahalla.region, "district": self.mahalla.district,
                    "gom": self.mahalla.gom,
                })
            if obj is None:
                self.warnings.append(
                    f"умумий r{row}: rahbar, telefon va tashkilot nomi bo'sh — "
                    f"obyekt qo'shilmadi")
                continue
            self.bump(f"Object {current_key}", created)
            seen_leaders.add((current_key, full_name))
            if phone != "-":
                seen_phones.add((current_key, phone))

            found = photo_index.get((current_key, full_name))
            if found:
                self.save_photo(obj, "avatar", *found)

        self._objects_from_sheets(ws, categories, seen_leaders, seen_phones)

    def _objects_from_sheets(self, ws, categories, seen, seen_phones):
        """
        `умумий` da yo'q obyektlarni alohida varaqlardan qo'shadi.

        Ba'zi fayllarda `умумий` obyektlar blokida ayrim turkumlar uchun faqat
        SONI yoziladi (rahbar/manzil ustunlari bo'sh), to'liq ma'lumot esa o'z
        varag'ida bo'ladi. Ular ham qo'shilmasa, obyektlar yo'qolib qoladi.
        """
        wb = ws.parent
        # Bazada turkumi yo'q varaqlar (kutubxona, ziyoratgoh, qabriston, meros)
        for sheet in wb.sheetnames:
            if norm_key(sheet) in CATEGORY_NOT_IN_DB:
                self.skipped_categories.add(sheet)

        for sheet, key in PHOTO_SHEET_CATEGORY.items():
            sheet = resolve_sheet(wb, sheet)
            if sheet is None:
                continue
            category = categories.get(key)
            if category is None:
                self.warnings.append(
                    f"{sheet}: ObjectCategory '{key}' bazada yo'q, o'tkazib yuborildi")
                continue

            s = wb[sheet]
            header_row, cols = find_header(s, OBJECT_HEADER)
            if not cols.get("leader"):
                self.warnings.append(f"{sheet}: sarlavha topilmadi, o'tkazib yuborildi")
                continue

            def cell(row, name):
                col = cols.get(name)
                return clean(s.cell(row, col).value) if col else ""

            for row in range(header_row + 1, s.max_row + 1):
                if not cell(row, "leader"):
                    continue
                full_name = person_name(s.cell(row, cols["leader"]).value)
                # `norm_name` bo'sh — katakda faqat tinish belgisi bor ("\", "-"):
                # ism emas, lekin `if not full_name` uni ushlamaydi.
                if not norm_name(full_name) or (key, full_name) in seen:
                    continue          # `умумий` dan allaqachon qo'shilgan

                # Bir obyekt `умумий` da tashkilot nomi bilan, o'z varag'ida esa
                # boshqacha ("Savdo majmuasi" / "Yo'q") yozilishi mumkin —
                # o'shanda ism bo'yicha topilmay, ikkinchi nusxa yaratilardi.
                phone = parse_phone(cell(row, "phone")) or "-"
                if phone != "-" and (key, phone) in seen_phones:
                    self.warnings.append(
                        f"{sheet} r{row}: '{full_name}' shu turkumda bir xil telefon "
                        f"bilan allaqachon bor ({phone}) — takror qo'shilmadi")
                    continue

                lat, lng, note = parse_location(cell(row, "location"))
                cx, cy = self.check_coords(lat, lng, f"{sheet} r{row}", note)

                obj, created = self.upsert_by_name(
                    Object, full_name, {"category": category},
                    defaults={
                        "phone_number": phone,
                        "address": to_latin(cell(row, "address")) or None,
                        "coordinate_x": cx, "coordinate_y": cy,
                        "region": self.mahalla.region,
                        "district": self.mahalla.district,
                        "gom": self.mahalla.gom,
                    })
                if obj is None:
                    continue
                self.bump(f"Object {key}", created)
                seen.add((key, full_name))
                if phone != "-":
                    seen_phones.add((key, phone))
                self.warnings.append(
                    f"{sheet} r{row}: '{full_name}' faqat shu varaqda bor "
                    f"(умумий da yo'q) — tashkilot nomisiz qo'shildi")
                self.save_photo(obj, "avatar", sheet, row)

    def _object_photo_index(self, ws):
        """
        {(ObjectCategory.key, rahbar F.I.Sh lotin): (varaq, qator)}

        Obyekt rasmlari alohida varaqlarda — rahbar ismi bo'yicha bog'lanadi.
        """
        index = {}
        wb = ws.parent
        for sheet, key in PHOTO_SHEET_CATEGORY.items():
            sheet = resolve_sheet(wb, sheet)
            if sheet is None:
                continue
            s = wb[sheet]
            # Sarlavha qatori fayldan faylga siljiydi — matn bo'yicha topiladi.
            header_row, cols = find_header(s, OBJECT_HEADER)
            leader_col = cols.get("leader")
            if not leader_col:
                continue
            for row in range(header_row + 1, s.max_row + 1):
                name = person_name(s.cell(row, leader_col).value)
                if name:
                    index[(key, name)] = (sheet, row)
        return index

    # -------------------------------------------------------------- utility
    @staticmethod
    def _find_row(ws, predicate, limit=None):
        for r in range(1, (limit or ws.max_row) + 1):
            try:
                if predicate(r):
                    return r
            except Exception:
                continue
        return None

    def report(self):
        self.stdout.write(self.style.MIGRATE_HEADING("\nNatija"))
        for key in sorted(self.stats):
            self.stdout.write(f"  {key:44s} {self.stats[key]}")
        if not self.stats:
            self.stdout.write("  (hech narsa qo'shilmadi)")

        self.stdout.write(self.style.MIGRATE_HEADING(
            f"\nO'tkazib yuborilgan varaqlar: {', '.join(SHEETS_SKIPPED)} "
            f"(mos model yo'q)"))
        if self.skipped_categories:
            self.stdout.write(self.style.WARNING(
                f"Bazada ObjectCategory yo'qligi uchun o'tkazildi: "
                f"{', '.join(sorted(self.skipped_categories))}"))

        if self.warnings:
            self.stdout.write(self.style.WARNING(
                f"\nOgohlantirishlar ({len(self.warnings)}):"))
            for w in self.warnings[:80]:
                self.stdout.write(f"  - {w}")
            if len(self.warnings) > 80:
                self.stdout.write(f"  … yana {len(self.warnings) - 80} ta")

        if self.dry:
            self.stdout.write(self.style.WARNING(
                "\nDRY-RUN edi — baza o'zgarmadi."))
        else:
            self.stdout.write(self.style.SUCCESS("\nYuklash yakunlandi."))
